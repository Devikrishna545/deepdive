"""User-uploaded documents: ingest (PDF / DOCX / TXT) and retrieve.

Vector store: embeddings from Mesh, stored in SQLite (doc_chunks.embedding_json) and
searched with numpy cosine similarity. Small, zero-ops, good enough for a handful of
documents. Swap for Chroma/FAISS behind the same two functions if needed.
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
from sqlalchemy import text as sa_text
from sqlmodel import select

from app.core.utils import clean_document_text, new_id
from app.db.models import DocChunk, Document
from app.db.session import session_scope
from app.llm.mesh_client import get_mesh
from app.tools.chunker import split_text

ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}
# Web pages are cut at MIN_CHUNK_CHARS to drop menus and footers; a person's own file is kept almost
# whole (a one-line note is still evidence they chose to give us).
DOC_MIN_CHUNK_CHARS = 20


class DocumentError(ValueError):
    """The file cannot be used: unreadable, or no text in it (e.g. a scanned PDF without a text layer)."""


def _read_pages(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import fitz

        with fitz.open(path) as doc:
            return [(i + 1, p.get_text()) for i, p in enumerate(doc)]
    if suffix == ".docx":
        import docx

        return [(None, "\n".join(p.text for p in docx.Document(str(path)).paragraphs))]
    return [(None, path.read_text(encoding="utf-8", errors="ignore"))]


def ingest_document(path: Path, filename: str) -> Document:
    try:
        pages = _read_pages(path)
    except Exception as exc:  # corrupt or mislabelled file: tell the person, never a 500
        raise DocumentError(f"Could not read {filename}: the file looks damaged or is not a real {path.suffix.lstrip('.').upper()}.") from exc
    pieces = [(page, text) for page, page_text in pages
              for text in split_text(clean_document_text(page_text), min_chars=DOC_MIN_CHUNK_CHARS)]
    if not pieces:
        raise DocumentError(f"{filename} has no readable text. Scanned PDFs need a text layer (run OCR first).")
    vectors = get_mesh().embed([t for _, t in pieces]) if pieces else []
    doc = Document(id=new_id("doc"), filename=filename, pages=len(pages))
    with session_scope() as s:
        s.add(doc)
        for (page, text), vec in zip(pieces, vectors, strict=True):
            s.add(DocChunk(id=new_id("dch"), document_id=doc.id, page=page, text=text, embedding_json=json.dumps(vec)))
        s.commit()
    return doc


def get_documents(document_ids: list[str]) -> dict[str, Document]:
    if not document_ids:
        return {}
    with session_scope() as s:
        docs = s.exec(select(Document).where(Document.id.in_(document_ids)))
        return {d.id: d for d in docs}


def front_chunks(document_ids: list[str], n: int = 2) -> list[DocChunk]:
    """The first `n` chunks of each document, in reading order: title, authors, abstract. Meaning
    search rarely ranks the title page for "who wrote it / when / where was it published"."""
    out: list[DocChunk] = []
    with session_scope() as s:
        for doc_id in document_ids:
            stmt = select(DocChunk).where(DocChunk.document_id == doc_id).order_by(sa_text("rowid")).limit(n)
            out += list(s.exec(stmt))
    return out


def document_preview(document_ids: list[str], chars: int = 700) -> str:
    """Each document's name and opening text, for the planner in document mode (untrusted text,
    framed as a SOURCE block like web text)."""
    docs = get_documents(document_ids)
    blocks = []
    for doc_id in document_ids:
        head = clean_document_text(" ".join(c.text for c in front_chunks([doc_id], 3)))[:chars]
        head = head.replace("<SOURCE", "&lt;SOURCE").replace("</SOURCE", "&lt;/SOURCE")
        name = docs[doc_id].filename if doc_id in docs else doc_id
        blocks.append(f"<SOURCE name={name}>\n{head}\n</SOURCE>")
    return "\n".join(blocks)


def retrieve(document_ids: list[str], query: str, k: int = 4) -> list[DocChunk]:
    if not document_ids:
        return []
    with session_scope() as s:
        rows = list(s.exec(select(DocChunk).where(DocChunk.document_id.in_(document_ids))))
    if not rows:
        return []
    q = np.asarray(get_mesh().embed([query])[0], dtype=float)
    m = np.asarray([json.loads(r.embedding_json) for r in rows], dtype=float)
    scores = (m @ q) / (np.linalg.norm(m, axis=1) * np.linalg.norm(q) + 1e-9)
    order = np.argsort(-scores)[:k]
    return [rows[i] for i in order]
